"""文档解析:.md/.txt 直读(编码兜底)、.pdf(pypdf)、.docx(python-docx)。

解析器只在 ingest 时运行,失败要把错误带回文档状态,所以这里只抛异常,
由 IngestWorker 捕获落库。解析依赖(pypdf/python-docx)懒加载:不上传对应
类型就不 import,保持启动轻量。
"""
from __future__ import annotations

from pathlib import Path


class UnsupportedFileType(ValueError):
    pass


def extract_text(path: Path, ext: str) -> str:
    ext = ext.lower().lstrip(".")
    if ext in ("md", "markdown", "txt"):
        return _read_text(path)
    if ext == "pdf":
        return _read_pdf(path)
    if ext == "docx":
        return _read_docx(path)
    raise UnsupportedFileType(f"不支持的文件类型: .{ext}")


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    for enc in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("解析 PDF 需要 pypdf 依赖（pip install pypdf）") from exc
    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return "\n\n".join(p for p in pages if p.strip())


def _read_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("解析 DOCX 需要 python-docx 依赖") from exc
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)
